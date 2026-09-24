#!/usr/bin/env python3
import argparse,json,re,sys
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook
T={'x','1','true','yes','có'}
def n(v):return re.sub(r'\s+',' ',str(v or '').strip()).casefold()
def m(v):return n(v) in T
def main():
 p=argparse.ArgumentParser(description='Create per-ID classification plan from MASTER_TRACKING.');p.add_argument('input');p.add_argument('--sheet',default='MASTER_TRACKING');p.add_argument('--output','-o');a=p.parse_args();path=Path(a.input)
 if not path.exists():return 3
 wb=load_workbook(path,data_only=True,read_only=True)
 if a.sheet not in wb.sheetnames:return 4
 ws=wb[a.sheet];heads=[ws.cell(3,c).value for c in range(1,ws.max_column+1)];rows=[dict(zip(heads,r)) for r in ws.iter_rows(min_row=4,values_only=True)];g=defaultdict(list)
 for excel_row,r in enumerate(rows,4):r['_excel_row']=excel_row;g[r['ID']].append(r)
 items=[]
 for key,h in sorted(g.items()):
  h.sort(key=lambda x:x['Snapshot']);r=h[-1];prog=n(r['Kết quả/tiến độ đến thời điểm báo cáo']);days=r['Số ngày còn lại'];pic=n(r.get('Chịu trách nhiệm (PIC)'));sig=[]
  if isinstance(days,(int,float)) and days<0:sig+=['overdue']
  if isinstance(days,(int,float)) and 0<=days<=14:sig+=['due_0_14']
  if days=='Không giao thời hạn':sig+=['no_specific_deadline']
  if isinstance(r.get('Chênh lệch dự kiến so với thời hạn (Ngày)'),(int,float)) and r['Chênh lệch dự kiến so với thời hạn (Ngày)']>0:sig+=['forecast_after_deadline']
  terms=['đang phối hợp','đang trình','chờ phê duyệt','dự kiến','trình lại','góp ý lần 2','không thay đổi','đã nhận chỉ đạo','thực hiện theo chỉ đạo']
  sig += ['activity_language:'+x for x in terms if x in prog]
  marks=[c for c in ['Hoàn Thành','Thực hiện đúng tiến độ','Chưa hoàn thành'] if m(r[c])]
  if len(marks)>1:sig+=['multiple_status_marks']
  if 'Thực hiện đúng tiến độ' in marks and 'overdue' in sig:sig+=['overdue_but_marked_on_track']
  if 'Hoàn Thành' in marks and any(x.startswith('activity_language:') for x in sig):sig+=['completion_needs_evidence_review']
  if len(h)>1 and prog==n(h[-2]['Kết quả/tiến độ đến thời điểm báo cáo']):sig+=['no_substantive_progress']
  facts={'output_identifiable':bool(n(r['Tóm tắt nội dung chỉ đạo'])),'timing_identifiable':days is not None,'progress_informative':bool(prog and prog not in {'đã nhận chỉ đạo','thực hiện theo chỉ đạo'}),'pic_identifiable':bool(pic)}
  if not facts['output_identifiable'] or (not facts['timing_identifiable'] and not facts['progress_informative'] and not facts['pic_identifiable']):cand='XÁM'
  elif set(sig)&{'overdue','forecast_after_deadline','overdue_but_marked_on_track','multiple_status_marks'}:cand='CAM'
  elif any(x.startswith('activity_language:') for x in sig) or set(sig)&{'no_specific_deadline','completion_needs_evidence_review'}:cand='VÀNG'
  else:cand='REVIEW_FOR_XANH_OR_VÀNG'
  items.append({'id':key,'latest_snapshot':str(r['Snapshot'].date()),'snapshot_count':len(h),'source_rows':[{'sheet':x['Source_Sheet'],'row':x['Source_Row'],'master_row':x['_excel_row']} for x in h],'minimum_facts':facts,'status_marks':marks,'observed_signals':sorted(set(sig)),'rule_based_candidate':cand,'red_assessment_required':True,'current':{k:v for k,v in r.items() if not k.startswith('_')}})
 out={'schema_version':'2.0','source_sheet':a.sheet,'warning':'Candidates are not final. ĐỎ requires severity plus authority assessment; XANH requires positive evidence.','items':items};text=json.dumps(out,ensure_ascii=False,indent=2,default=str)
 if a.output:Path(a.output).write_text(text+'\n',encoding='utf-8');print(json.dumps({'status':'ok','items':len(items)},ensure_ascii=False))
 else:print(text)
 return 0
if __name__=='__main__':raise SystemExit(main())
