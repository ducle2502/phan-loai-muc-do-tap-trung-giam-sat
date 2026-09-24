#!/usr/bin/env python3
import argparse,json,sys,re
from collections import Counter,defaultdict
from pathlib import Path
from openpyxl import load_workbook
TRUTH={'x','1','true','yes','có'}
def marked(v):return str(v).strip().lower() in TRUTH if v is not None else False
def main():
 p=argparse.ArgumentParser(description='Profile MASTER_TRACKING without modifying the workbook.');p.add_argument('input');p.add_argument('--sheet',default='MASTER_TRACKING');p.add_argument('--output','-o');a=p.parse_args();path=Path(a.input)
 if not path.exists():return 3
 wb=load_workbook(path,data_only=True,read_only=True)
 if a.sheet not in wb.sheetnames:print('ERROR sheet_not_found',file=sys.stderr);return 4
 ws=wb[a.sheet];headers=[ws.cell(3,c).value for c in range(1,ws.max_column+1)];rows=[dict(zip(headers,r)) for r in ws.iter_rows(min_row=4,values_only=True)]
 required={'Snapshot','ID','Source_Sheet','Source_Row','Số ngày còn lại','Kết quả/tiến độ đến thời điểm báo cáo','Hoàn Thành','Thực hiện đúng tiến độ','Chưa hoàn thành'};missing=sorted(required-set(headers))
 if missing:print('ERROR missing_columns: '+', '.join(missing),file=sys.stderr);return 5
 by=defaultdict(list)
 for r in rows:by[r['ID']].append(r)
 unchanged=0
 for h in by.values():
  h.sort(key=lambda x:x['Snapshot'])
  if len(h)>1:
   n=lambda v:re.sub(r'\s+',' ',str(v or '').strip().casefold())
   unchanged+=n(h[-1]['Kết quả/tiến độ đến thời điểm báo cáo'])==n(h[-2]['Kết quả/tiến độ đến thời điểm báo cáo'])
 out={'sheet':a.sheet,'header_row':3,'data_start_row':4,'columns':headers,'row_count':len(rows),'stable_id_count':len(by),'snapshot_counts':Counter(str(r['Snapshot'].date()) for r in rows),'ids_with_multiple_snapshots':sum(len(h)>1 for h in by.values()),'unchanged_latest_progress_count':unchanged,'overdue_count':sum(isinstance(r['Số ngày còn lại'],(int,float)) and r['Số ngày còn lại']<0 for r in rows),'no_specific_deadline_count':sum(r['Số ngày còn lại']=='Không giao thời hạn' for r in rows),'due_0_14_count':sum(isinstance(r['Số ngày còn lại'],(int,float)) and 0<=r['Số ngày còn lại']<=14 for r in rows),'status_mark_counts':{'completed':sum(marked(r['Hoàn Thành']) for r in rows),'on_track':sum(marked(r['Thực hiện đúng tiến độ']) for r in rows),'incomplete':sum(marked(r['Chưa hoàn thành']) for r in rows),'multiple_marks':sum(sum(marked(r[c]) for c in ['Hoàn Thành','Thực hiện đúng tiến độ','Chưa hoàn thành'])>1 for r in rows)}}
 text=json.dumps(out,ensure_ascii=False,indent=2,default=dict)
 if a.output:Path(a.output).write_text(text+'\n',encoding='utf-8');print(json.dumps({'status':'ok','output':a.output},ensure_ascii=False))
 else:print(text)
 return 0
if __name__=='__main__':raise SystemExit(main())
